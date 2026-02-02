"""
Output generators for translation review and localized taxonomy files.
"""

import logging
import os
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from .config import TAXONOMY_FILES, LanguageConfig


class ReviewExcelGenerator:
    """Generate Excel files for translation review with side-by-side English and translated columns."""

    ENCODING = 'utf-8-sig'

    def __init__(self, lang_code: str, output_dir: str):
        """
        Initialize the review Excel generator.

        Args:
            lang_code: Target language code (e.g., 'sw', 'es')
            output_dir: Directory for output files
        """
        self.lang_code = lang_code.lower()
        self.lang_suffix = lang_code.upper()
        self.output_dir = output_dir
        self.lang_info = LanguageConfig.get_language_info(lang_code)

    def generate(self, translated_csv: str, file_config: dict[str, Any]) -> str | None:
        """
        Generate review Excel from translated CSV.

        Args:
            translated_csv: Path to translated CSV file
            file_config: Configuration for the taxonomy file

        Returns:
            Path to generated Excel file, or None on failure
        """
        if not Path(translated_csv).exists():
            logging.error(f"Translated CSV not found: {translated_csv}")
            return None

        try:
            df = pd.read_csv(translated_csv, encoding=self.ENCODING)
        except Exception:
            df = pd.read_csv(translated_csv, encoding='utf-8')

        # Build review dataframe with side-by-side columns
        review_columns = self._build_review_columns(df, file_config)
        review_df = df[review_columns].copy()

        # Generate output path
        base_name = Path(translated_csv).name.replace('.csv', '').replace('_translated_', '_review_')
        output_path = str(Path(self.output_dir) / f"{base_name}.xlsx")

        # Write Excel with formatting
        self._write_excel(review_df, output_path, file_config)

        logging.info(f"Generated review Excel: {Path(output_path).name}")
        return output_path

    def _build_review_columns(self, df: pd.DataFrame, file_config: dict[str, Any]) -> list[str]:
        """Build list of columns for review, alternating English and translated."""
        columns = []
        id_field = file_config['id_field']
        translation_fields = file_config['translation_fields']

        # Always include ID first
        if id_field in df.columns:
            columns.append(id_field)

        # Add CODE if present
        if 'CODE' in df.columns:
            columns.append('CODE')

        # Add side-by-side columns for each translatable field
        for field in translation_fields:
            if field in df.columns:
                columns.append(field)  # English
                translated_col = f"{field}_{self.lang_suffix}"
                if translated_col in df.columns:
                    columns.append(translated_col)  # Translated

        return columns

    def _write_excel(self, df: pd.DataFrame, output_path: str, file_config: dict[str, Any]) -> None:
        """Write DataFrame to Excel with formatting."""
        lang_name = self.lang_info['name'] if self.lang_info else self.lang_code.upper()

        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            # Write review data
            df.to_excel(writer, sheet_name='Review', index=False)

            # Access workbook for formatting
            workbook = writer.book
            worksheet = writer.sheets['Review']

            # Format headers
            from openpyxl.styles import Font, PatternFill, Alignment

            header_fill_en = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')  # Oxford Blue
            header_fill_translated = PatternFill(start_color='00B894', end_color='00B894', fill_type='solid')  # Tabiya Green
            header_font = Font(bold=True, color='FFFFFF')

            for col_idx, col_name in enumerate(df.columns, 1):
                cell = worksheet.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

                # Color based on whether it's English or translated
                if col_name.endswith(f'_{self.lang_suffix}'):
                    cell.fill = header_fill_translated
                else:
                    cell.fill = header_fill_en

            # Set column widths
            for col_idx, col_name in enumerate(df.columns, 1):
                if 'ALTLABELS' in col_name or 'DESCRIPTION' in col_name or 'DEFINITION' in col_name:
                    worksheet.column_dimensions[worksheet.cell(row=1, column=col_idx).column_letter].width = 50
                elif 'PREFERREDLABEL' in col_name:
                    worksheet.column_dimensions[worksheet.cell(row=1, column=col_idx).column_letter].width = 35
                else:
                    worksheet.column_dimensions[worksheet.cell(row=1, column=col_idx).column_letter].width = 20

            # Freeze header row
            worksheet.freeze_panes = 'A2'

            # Add stats sheet
            stats_data = self._calculate_stats(df, file_config)
            stats_df = pd.DataFrame(stats_data)
            stats_df.to_excel(writer, sheet_name='Stats', index=False)

    def _calculate_stats(self, df: pd.DataFrame, file_config: dict[str, Any]) -> list[dict]:
        """Calculate translation coverage statistics."""
        stats = []
        translation_fields = file_config['translation_fields']

        for field in translation_fields:
            translated_col = f"{field}_{self.lang_suffix}"
            if field in df.columns and translated_col in df.columns:
                total = len(df)
                english_filled = df[field].notna().sum()
                translated_filled = df[translated_col].apply(
                    lambda x: pd.notna(x) and str(x).strip() not in ['', '[EMPTY]']
                ).sum()

                stats.append({
                    'Field': field,
                    'Total Records': total,
                    'English Filled': english_filled,
                    'Translated Filled': translated_filled,
                    'Coverage %': round(translated_filled / total * 100, 1) if total > 0 else 0,
                })

        return stats


class LocalizedTaxonomyGenerator:
    """Generate localized taxonomy CSV files matching source structure."""

    ENCODING = 'utf-8-sig'

    def __init__(self, lang_code: str, output_dir: str, source_dir: str):
        """
        Initialize the localized taxonomy generator.

        Args:
            lang_code: Target language code (e.g., 'sw', 'es')
            output_dir: Directory for output files
            source_dir: Directory containing source taxonomy files
        """
        self.lang_code = lang_code.lower()
        self.lang_suffix = lang_code.upper()
        self.output_dir = str(Path(output_dir))
        self.source_dir = str(Path(source_dir))
        self.localized_dir = str(Path(output_dir) / 'localized')
        Path(self.localized_dir).mkdir(parents=True, exist_ok=True)

    def generate(self, translated_csv: str, file_config: dict[str, Any], file_name: str) -> str | None:
        """
        Generate localized taxonomy CSV from translated CSV.

        Args:
            translated_csv: Path to translated CSV file
            file_config: Configuration for the taxonomy file
            file_name: Original source file name (e.g., 'occupations.csv')

        Returns:
            Path to generated localized CSV, or None on failure
        """
        if not Path(translated_csv).exists():
            logging.error(f"Translated CSV not found: {translated_csv}")
            return None

        try:
            df = pd.read_csv(translated_csv, encoding=self.ENCODING)
        except Exception:
            df = pd.read_csv(translated_csv, encoding='utf-8')

        # Load source file to get original column order
        source_path = Path(self.source_dir) / file_name
        if not source_path.exists():
            logging.warning(f"Source file not found: {source_path}")
            return None

        try:
            source_df = pd.read_csv(str(source_path), encoding=self.ENCODING)
        except Exception:
            source_df = pd.read_csv(str(source_path), encoding='utf-8')

        # Build localized dataframe
        localized_df = self._build_localized_df(df, source_df, file_config)

        # Output path
        output_path = str(Path(self.localized_dir) / file_name)
        localized_df.to_csv(output_path, index=False, encoding=self.ENCODING)

        logging.info(f"Generated localized taxonomy: {file_name}")
        return output_path

    def _build_localized_df(
        self,
        translated_df: pd.DataFrame,
        source_df: pd.DataFrame,
        file_config: dict[str, Any],
    ) -> pd.DataFrame:
        """
        Build localized dataframe by replacing English with translated content.

        Args:
            translated_df: DataFrame with translations
            source_df: Original source DataFrame
            file_config: Configuration for the taxonomy file

        Returns:
            Localized DataFrame
        """
        id_field = file_config['id_field']
        translation_fields = file_config['translation_fields']

        # Start with source structure
        localized_df = source_df.copy()

        # Build lookup from translated data
        translation_lookup = {}
        for _, row in translated_df.iterrows():
            record_id = row.get(id_field)
            if pd.notna(record_id):
                translation_lookup[record_id] = row

        # Replace fields with translations where available
        for idx, row in localized_df.iterrows():
            record_id = row.get(id_field)
            if record_id in translation_lookup:
                translated_row = translation_lookup[record_id]

                for field in translation_fields:
                    translated_col = f"{field}_{self.lang_suffix}"
                    if translated_col in translated_row.index:
                        translated_value = translated_row[translated_col]
                        # Only replace if translation exists and is not empty/error
                        if pd.notna(translated_value) and str(translated_value).strip() not in ['', '[EMPTY]']:
                            if 'ERROR' not in str(translated_value):
                                localized_df.at[idx, field] = translated_value

        return localized_df

    def copy_unchanged_files(self) -> list[str]:
        """
        Copy taxonomy files that don't need translation (hierarchies, relations).

        Returns:
            List of copied file paths
        """
        unchanged_files = [
            'occupation_hierarchy.csv',
            'skill_hierarchy.csv',
            'occupation_to_skill_relations.csv',
            'skill_to_skill_relations.csv',
            'model_info.csv',
            'VERSION.txt',
        ]

        copied = []
        for file_name in unchanged_files:
            source_path = Path(self.source_dir) / file_name
            dest_path = Path(self.localized_dir) / file_name

            if source_path.exists():
                try:
                    if file_name.endswith('.csv'):
                        df = pd.read_csv(str(source_path), encoding=self.ENCODING)
                        df.to_csv(str(dest_path), index=False, encoding=self.ENCODING)
                    else:
                        import shutil
                        shutil.copy2(source_path, dest_path)
                    copied.append(file_name)
                    logging.info(f"Copied unchanged file: {file_name}")
                except Exception as e:
                    logging.warning(f"Could not copy {file_name}: {e}")

        return copied
